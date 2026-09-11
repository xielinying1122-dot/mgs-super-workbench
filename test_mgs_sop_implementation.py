import os
import time
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import mgs_dashboard as dashboard
import margys_daily_report as report


def test_dashboard_normalize_filters_totals_and_zero_spend(tmp_path):
    frame = pd.DataFrame(
        [
            ["合计", "全站", "合计", "n-total", 999, 1, 1, 1, 999],
            ["2026-09-01", "全站", "计划A", "n1", 0, 10, 1, 3, 20],
            ["2026-09-01", "全站", "计划A", "n1", 100, 10, 1, 3, 20],
        ],
        columns=["时间", "投放类型", "计划名称", "笔记ID", "消费", "展现量", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
    )
    result = dashboard.normalize(frame, tmp_path / "report.xlsx", "daily", "日报数据底表")
    assert len(result) == 1
    assert result.iloc[0]["spend"] == 100
    assert result.iloc[0]["point"] == "全站"
    assert result.iloc[0]["source_sheet"] == "日报数据底表"


def test_dashboard_normalize_removes_total_anywhere_and_standardizes_point(tmp_path):
    frame = pd.DataFrame(
        [
            ["2026-09-01", "计划A", "笔记A", "信息流-智能", 100, 10, 1, 2, 50],
            ["2026-09-01", "计划A", "合计汇总", "信息流", 999, 10, 1, 2, 50],
        ],
        columns=["时间", "计划名称", "笔记ID", "投放类型", "消费", "展现量", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
    )
    result = dashboard.normalize(frame, tmp_path / "report.xlsx", "daily")
    assert len(result) == 1
    assert result.iloc[0]["point"] == "信息流"


def test_creator_mapping_overrides_source_creator_name(tmp_path):
    frame = pd.DataFrame(
        [["2026-09-01", "计划A", "n1", "源文件昵称", "全站", 100, 10, 1, 2, 50]],
        columns=["时间", "计划名称", "笔记ID", "达人昵称", "投放类型", "消费", "展现量", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
    )
    result = dashboard.normalize(frame, tmp_path / "report.xlsx", "daily", creator_map={"n1": "飞书昵称"})
    assert result.iloc[0]["creator_name"] == "飞书昵称"


def test_authoritative_creator_mapping_is_not_overwritten(tmp_path):
    source = tmp_path / "daily.xlsx"
    source.touch()
    frame = pd.DataFrame(
        [["2026-09-01", "计划A", "n1", "原始昵称", "全站", 100, 10, 1, 2, 50]],
        columns=["时间", "计划名称", "笔记ID", "达人昵称", "投放类型", "消费", "展现量", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
    )
    sources = dashboard.read_sources
    original = dashboard._read_file_frames
    try:
        dashboard._read_file_frames = lambda _: ([("日报数据底表", frame)], {"n1": "匹配表昵称"})
        result = sources(tmp_path)["daily"]
    finally:
        dashboard._read_file_frames = original
    assert result.iloc[0]["creator_name"] == "匹配表昵称"


def test_dashboard_dedupe_keeps_newest_file(tmp_path):
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    old.touch()
    time.sleep(0.01)
    new.touch()
    base = {"date": pd.Timestamp("2026-09-01"), "point": "全站", "campaign": "A", "unit": "", "note_id": "n1", "targeting": "", "keyword": "", "source_mtime": 0, "source_file": "old.xlsx", "spend": 100}
    newer = dict(base, source_mtime=new.stat().st_mtime, source_file="new.xlsx", spend=140)
    older = dict(base, source_mtime=old.stat().st_mtime)
    result = dashboard._dedupe(pd.DataFrame([newer, older]))
    assert len(result) == 1
    assert result.iloc[0]["source_file"] == "new.xlsx"
    assert result.iloc[0]["spend"] == 140


def test_dashboard_does_not_fallback_to_legacy_file_when_input_is_empty(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    legacy = tmp_path / "legacy.xlsx"
    legacy.touch()
    frame = pd.DataFrame(
        [["2026-09-01", "全站", "计划A", "n1", 100, 10, 1, 2, 50]],
        columns=["时间", "投放类型", "计划名称", "笔记ID", "消费", "展现量", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
    )
    monkeypatch.setattr(dashboard, "FALLBACK", legacy)
    monkeypatch.setattr(dashboard, "_read_file_frames", lambda _path: ([('日报数据底表', frame)], {}))

    result = dashboard.read_sources(input_dir)

    assert all(table.empty for table in result.values())


def test_action_id_is_repeatable_and_separates_points():
    asof = pd.Timestamp("2026-09-04")
    full_site = dashboard._stable_action_id(asof, "笔记", "全站", "note-1")
    feed = dashboard._stable_action_id(asof, "笔记", "信息流", "note-1")
    assert full_site == dashboard._stable_action_id(asof, "笔记", "全站", "note-1")
    assert full_site != feed
    assert full_site.endswith("-" + dashboard.hashlib.sha1("笔记|全站|note-1".encode("utf-8")).hexdigest()[:10])


def test_candidates_have_unique_action_ids_when_legacy_ids_collide():
    rows = []
    for point, object_id in [("全站", "note-1"), ("信息流", "note-1"), ("搜索", "note-1")]:
        rows.append({
            "date": pd.Timestamp("2026-09-01"), "point": point, "campaign": "计划-" + point,
            "unit": "", "note_id": object_id, "note_name": "笔记", "creator_name": "达人",
            "targeting": "", "keyword": "", "spend": 100, "impressions": 1000, "clicks": 20,
            "interactions": 0, "visits": 2, "gmv": 200, "gsv": None, "gmv_available": True,
            "data_status": "可判断", "source_kind": "daily", "report_type": "笔记/创意",
            "source_file": "a.xlsx", "source_sheet": "日报", "source_path": "a.xlsx",
            "source_mtime": 1, "gmv_source_field": "GMV", "visits_source_field": "进店",
        })
    sources = {"daily": pd.DataFrame(rows), "targeting": pd.DataFrame(), "keyword": pd.DataFrame()}
    candidates = dashboard._all_candidates(sources, pd.Timestamp("2026-09-04"), {})
    ids = [item["action_id"] for item in candidates]
    assert len(ids) == len(set(ids))
    assert {item["point"] for item in candidates if item["level"] == "笔记"} == {"全站", "信息流", "搜索"}


def test_execution_migration_maps_unique_legacy_id_and_drops_blank_duplicates():
    candidates = [
        {"action_id": "MGS-20260904-计划-全站-计划A-deadbeef01", "__old_action_id": "MGS-20260903-计划-计划A"},
    ]
    old_execution = [
        {"动作编号": "MGS-20260903-计划-计划A", "状态": "待确认", "复核备注": "待人工确认后填写执行信息"},
        {"动作编号": "MGS-20260903-计划-计划A", "状态": "已执行", "执行时间": "2026-09-04 10:00", "复核备注": "已完成"},
    ]
    migrated = dashboard._migrate_execution_records(old_execution, {}, candidates)
    assert len(migrated) == 1
    assert migrated[0]["动作编号"] == candidates[0]["action_id"]
    assert migrated[0]["状态"] == "已执行"


def test_execution_migration_archives_filled_ambiguous_legacy_id():
    old_id = "MGS-20260903-笔记-截断摘要"
    candidates = [
        {"action_id": "MGS-20260904-笔记-全站-note-a-aaaaaaaaaa", "__old_action_id": old_id},
        {"action_id": "MGS-20260904-笔记-信息流-note-a-bbbbbbbbbb", "__old_action_id": old_id},
    ]
    old_execution = [
        {"动作编号": old_id, "状态": "待确认", "复核备注": "待人工确认后填写执行信息"},
        {"动作编号": old_id, "状态": "已执行", "执行时间": "2026-09-04 10:00", "复核备注": "人工填写"},
    ]
    migrated = dashboard._migrate_execution_records(old_execution, {}, candidates)
    assert len(migrated) == 1
    assert migrated[0]["动作编号"].startswith("历史-" + old_id + "-")
    assert "需人工重新关联" in migrated[0]["复核备注"]


def test_dashboard_decision_respects_100_yuan_floor():
    stats = {"spend": 80, "roi": 0.1, "gmv_available": True}
    recent = {"spend": 80, "roi": 0.1}
    previous = {"spend": 80, "roi": 0.2}
    decision = dashboard._decision(stats, recent, previous, 3, pd.Timestamp("2026-09-04"), pd.Timestamp("2026-09-06"))
    assert decision["action"] != "停投"


def test_dashboard_decision_holds_until_attribution_matures():
    stats = {"spend": 200, "roi": 0.1, "gmv_available": True}
    recent = {"spend": 100, "roi": 0.1}
    previous = {"spend": 100, "roi": 0.2}
    decision = dashboard._decision(stats, recent, previous, 3, pd.Timestamp("2026-09-04"), pd.Timestamp("2026-09-04"))
    assert decision["action"] == "待核验"
    assert "未成熟" in decision["data_status"]


def test_dashboard_decision_prefers_recent_improvement_over_two_low_periods():
    stats = {"spend": 300, "roi": 1.0, "gmv_available": True}
    recent = {"spend": 150, "roi": 1.2}
    previous = {"spend": 150, "roi": 0.8}
    decision = dashboard._decision(stats, recent, previous, 4, pd.Timestamp("2026-09-04"), pd.Timestamp("2026-09-06"))
    assert decision["action"] == "降预算观察"


def test_dashboard_decision_observes_when_low_roi_lacks_two_period_comparison():
    stats = {"spend": 300, "roi": 1.0, "gmv_available": True}
    recent = {"spend": 150, "roi": None}
    previous = {"spend": 150, "roi": 0.8}
    decision = dashboard._decision(stats, recent, previous, 4, pd.Timestamp("2026-09-04"), pd.Timestamp("2026-09-06"))
    assert decision["action"] == "观察"
    assert "完整的连续周期对比" in decision["reason"]


def test_dashboard_build_uses_daily_date_asof_when_other_source_is_newer(tmp_path):
    daily = pd.DataFrame(
        [{"date": pd.Timestamp("2026-09-04"), "point": "全站", "campaign": "A", "unit": "", "note_id": "n1", "note_name": "N1", "creator_name": "C1", "targeting": "", "keyword": "", "spend": 100, "impressions": 10, "clicks": 2, "interactions": 0, "visits": 5, "gmv": 200, "gsv": None, "gmv_available": True, "data_status": "可判断", "source_kind": "daily", "report_type": "笔记/创意", "source_file": "a.xlsx", "source_sheet": "日报数据底表", "source_path": "a.xlsx", "source_mtime": 1, "gmv_source_field": "GMV", "visits_source_field": "进店"}]
    )
    targeting = daily.copy()
    targeting["date"] = pd.Timestamp("2026-09-06")
    targeting["source_kind"] = "targeting"
    targeting["report_type"] = "定向"
    sources = {"daily": daily, "targeting": targeting, "keyword": pd.DataFrame()}
    output = tmp_path / "mgs.xlsx"
    dashboard.build_workbook(sources, output, tmp_path)
    wb = load_workbook(output, data_only=False)
    assert "2026-09-01至2026-09-04" in wb["日报汇总"]["A1"].value
    wb.close()


def test_dashboard_refuses_to_overwrite_locked_workbook(tmp_path):
    output = tmp_path / "mgs.xlsx"
    output.write_bytes(b"existing")
    (tmp_path / ".~lock.mgs.xlsx#").write_text("locked", encoding="utf-8")
    with pytest.raises(RuntimeError, match="正在被 WPS/Excel 占用"):
        dashboard.build_workbook({"daily": pd.DataFrame(), "targeting": pd.DataFrame(), "keyword": pd.DataFrame()}, output, tmp_path)


def test_daily_report_update_requires_date_and_metrics_change():
    previous = {"latest_date": "2026-09-05", "metrics": {"整体": {"spend": 100}}}
    same_date_changed_metrics = {"latest_date": "2026-09-05", "metrics": {"整体": {"spend": 120}}}
    changed_date_changed_metrics = {"latest_date": "2026-09-06", "metrics": {"整体": {"spend": 120}}}
    assert not report._is_updated(same_date_changed_metrics, previous)
    assert report._is_updated(changed_date_changed_metrics, previous)


def test_daily_report_uses_tot_window_and_gsv_is_not_invented():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日)"]]
    rows.extend([
        ["46266", "100", "1000", "10", "5", "200"],
        ["46267", "50", "500", "5", "2", "100"],
    ])
    workbook = {"Byday数据": rows, "CID底表": []}
    output = report.build_report(workbook)
    assert "9月1日-9月2日累计" in output
    assert "消耗：150.00" in output
    assert "去退后成交金额（GSV-实时更新截止现在）：待补录" in output


def test_daily_report_marks_partial_gsv_as_pending_and_uses_plain_numbers():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）", "去退后成交金额"]]
    rows.extend([
        ["46266", "1234.5", "1000", "10", "5", "2000", "1800"],
        ["46267", "678.9", "500", "5", "2", "1000", ""],
    ])
    output = report.build_report({"Byday数据": rows, "CID底表": []})
    assert "消耗：1913.40" in output
    assert "去退后成交金额（GSV-实时更新截止现在）：待补录" in output


def test_daily_report_missing_gmv_column_is_pending_not_zero():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)"]]
    rows.append(["46266", "100", "1000", "10", "5"])
    parsed = report.parse_daily(rows)
    output = report.build_report({"Byday数据": rows, "CID底表": []})
    assert parsed[0]["gmv"] is None
    assert parsed[0]["gmv_available"] is False
    assert "店铺成交GMV（30日)：待补录" in output
    assert "店铺成交ROI（30日)：待补录" in output


def test_daily_report_reads_targeting_and_keyword_detail_sheets():
    byday = [
        ["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）", "投放类型"],
        ["46266", "100", "1000", "10", "5", "200", "全站"],
    ]
    cid = [
        ["时间", "计划名称", "笔记ID", "创意名称", "消费", "展现量", "点击量", "行业商品进店量30日", "行业商品GMV（30日）", "投放类型"],
        ["46266", "计划A", "note-1", "创意A", "100", "1000", "10", "5", "200", "全站"],
    ]
    targeting = [
        ["时间", "精准定向", "计划名称", "消费", "行业商品进店量30日", "行业商品GMV（30日）"],
        ["46266", "高意向人群", "计划A", "100", "5", "200"],
    ]
    keyword = [
        ["时间", "关键词", "计划名称", "消费", "行业商品进店量30日", "行业商品GMV（30日）"],
        ["46266", "胶原面膜", "计划A", "100", "5", "200"],
    ]
    workbook = {"Byday数据": byday, "CID底表": cid, "人群包-底表": targeting, "关键词-底表": keyword}
    output = report.build_report(workbook)
    assert "定向层" in output
    assert "高意向人群" in output
    assert "笔记层" in output
    assert "note-1" in output
    assert "关键词层" in output
    assert "胶原面膜" in output
    assert "当前待补数据层级：" not in output


def test_daily_report_marks_missing_dimension_sheet_as_pending():
    byday = [
        ["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"],
        ["46266", "100", "1000", "10", "5", "200"],
    ]
    output = report.build_report({"Byday数据": byday, "CID底表": []})
    assert "当前待补数据层级：计划/笔记、定向、关键词" in output
    assert "定向层：未读取到可用明细，待补数据。" in output
    assert "关键词层：未读取到可用明细，待补数据。" in output


def test_daily_report_ignores_zero_spend_placeholder_rows():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"]]
    rows.extend([
        ["46266", "100", "1000", "10", "5", "200"],
        ["46266", "0", "0", "0", "0", ""],
    ])
    daily = report.parse_daily(rows)
    summary = report.aggregate(daily)
    assert len(daily) == 1
    assert summary["gmv_available"] is True
    assert summary["roi"] == 2.0


def test_daily_report_ignores_total_rows_marked_in_object_column():
    rows = [["时间", "计划名称", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"]]
    rows.extend([
        ["46266", "计划A", "100", "1000", "10", "5", "200"],
        ["46266", "合计", "999", "999", "999", "999", "999"],
    ])
    parsed = report.parse_cid(rows)
    assert len(parsed) == 1
    assert parsed[0]["spend"] == 100


def test_daily_report_partial_gmv_does_not_turn_missing_rows_into_zero():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"]]
    rows.extend([
        ["46266", "100", "1000", "10", "5", "200"],
        ["46267", "50", "500", "5", "2", ""],
    ])
    daily = report.parse_daily(rows)
    summary = report.aggregate(daily)
    assert summary["gmv_available"] is False
    assert summary["roi"] is None


def test_cid_missing_gmv_keeps_spend_rows_for_pending_actions():
    rows = [["时间", "计划名称", "消费", "展现量", "点击量", "行业商品进店量30日"]]
    rows.append(["46266", "计划A", "100", "1000", "10", "5"])
    parsed = report.parse_cid(rows)
    summary = report.aggregate(parsed)
    assert len(parsed) == 1
    assert summary["spend"] == 100
    assert summary["roi"] is None


def test_daily_report_update_signature_contains_placement_metrics():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日)"]]
    rows.append(["46266", "100", "1000", "10", "5", "200"])
    daily = report.parse_daily(rows)
    signature = report._update_signature({}, daily, [])
    assert signature["latest_date"] == "2026-09-01"
    assert set(signature["metrics"]) == {"整体", "全站", "信息流", "搜索"}


def test_daily_report_signature_uses_byday_placement_rows_not_cid_rows():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）", "投放类型"]]
    rows.extend([
        ["46266", "100", "1000", "10", "5", "200", "全站"],
        ["46267", "50", "500", "5", "2", "100", "信息流"],
    ])
    daily = report.parse_daily(rows)
    cid_rows = [{"point": "搜索", "spend": 999, "impressions": 999, "clicks": 999, "visits": 999, "gmv": 999}]
    signature = report._update_signature({}, daily, cid_rows)
    assert signature["point_data_available"] is True
    assert signature["metrics"]["整体"]["spend"] == 150
    assert signature["metrics"]["全站"]["spend"] == 100
    assert signature["metrics"]["信息流"]["spend"] == 50
    assert signature["metrics"]["搜索"] is None


def test_daily_report_marks_missing_byday_placement_data():
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"]]
    rows.append(["46266", "100", "1000", "10", "5", "200"])
    daily = report.parse_daily(rows)
    signature = report._update_signature({}, daily, [])
    assert signature["point_data_available"] is False
    assert signature["metrics"]["全站"] is None


def test_no_send_run_still_persists_update_state(tmp_path, monkeypatch):
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"]]
    rows.append(["46266", "100", "1000", "10", "5", "200"])
    monkeypatch.setattr(report, "REPORT_DIR", tmp_path)
    monkeypatch.setattr(report, "STATE_PATH", tmp_path / "state.json")
    monkeypatch.setattr(report, "read_workbook", lambda: {"Byday数据": rows, "CID底表": []})
    monkeypatch.setattr(report.sys, "argv", ["margys_daily_report.py", "--no-send", "--force"])
    report.main()
    assert (tmp_path / "state.json").exists()
    state = (tmp_path / "state.json").read_text(encoding="utf-8")
    assert "2026-09-01" in state


def test_retry_baseline_uses_current_snapshot_when_date_changes_without_metrics(tmp_path, monkeypatch):
    monkeypatch.setattr(report, "REPORT_DIR", tmp_path)
    monkeypatch.setattr(report, "STATE_PATH", tmp_path / "state.json")
    previous = {
        "signature": {
            "latest_date": "2026-09-04",
            "point_data_available": False,
            "metrics": {"整体": {"spend": 100}},
        }
    }
    report._save_state({**previous, "retry_pending": False})
    current = {
        "latest_date": "2026-09-05",
        "point_data_available": False,
        "metrics": {"整体": {"spend": 100}},
    }
    assert not report._is_updated(current, previous["signature"])
    report._save_state(report._formal_state(current, pending_retry=True))
    state = report._load_state()
    assert state["retry_pending"] is True
    assert state["signature"] == current


def test_daily_report_reminder_lists_missing_detail_layers(tmp_path):
    rows = [["时间", "消费", "展现量", "点击量", "店铺访问量(15日)", "行业商品GMV（30日）"]]
    rows.append(["46266", "100", "1000", "10", "5", "200"])

    reminder = report._reminder(
        {"Byday数据": rows, "CID底表": []},
        "整体和点位关键指标未发生变化",
    )

    assert "飞书日报最新可用数据日：2026-09-01" in reminder
    assert "当前待补数据层级：计划/笔记、定向、关键词" in reminder


def test_workbook_preserves_execution_records(tmp_path):
    output = tmp_path / "mgs.xlsx"
    source = pd.DataFrame(
        [{"date": pd.Timestamp("2026-09-01"), "point": "全站", "campaign": "A", "unit": "", "note_id": "n1", "note_name": "N1", "creator_name": "C1", "targeting": "", "keyword": "", "spend": 100, "impressions": 10, "clicks": 2, "interactions": 0, "visits": 5, "gmv": 200, "gsv": None, "gmv_available": True, "data_status": "可判断", "source_kind": "daily", "report_type": "笔记/创意", "source_file": "a.xlsx", "source_sheet": "日报数据底表", "source_path": "a.xlsx", "source_mtime": 1, "gmv_source_field": "GMV", "visits_source_field": "进店"}]
    )
    sources = {"daily": source, "targeting": pd.DataFrame(), "keyword": pd.DataFrame()}
    dashboard.build_workbook(sources, output, tmp_path)
    wb = load_workbook(output)
    ws = wb["执行复核"]
    ws["G5"] = "已执行"
    ws["B5"] = "2026-09-02 10:00"
    wb.save(output)
    dashboard.build_workbook(sources, output, tmp_path)
    wb = load_workbook(output, data_only=False)
    assert wb["执行复核"]["G5"].value == "已执行"
    assert wb["执行复核"]["B5"].value == "2026-09-02 10:00"
