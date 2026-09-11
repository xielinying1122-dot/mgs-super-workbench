import json
from pathlib import Path

import pandas as pd
from fastapi.testclient import TestClient

from mgs_workbench_app import create_app


def _write_creative(path: Path, creator: str = "达人甲", spend: float = 120) -> None:
    pd.DataFrame(
        [
            {
                "时间": "2026-09-01",
                "投放类型": "信息流",
                "计划名称": "计划A",
                "笔记ID": "note-1",
                "创意名称": "笔记一",
                "达人": creator,
                "消费": spend,
                "展现量": 2000,
                "点击量": 100,
                "互动量": 40,
                "行业商品进店量30日": 25,
                "行业商品GMV（30日)": 360,
            },
            {
                "时间": "2026-09-02",
                "投放类型": "搜索",
                "计划名称": "计划B",
                "笔记ID": "note-2",
                "创意名称": "笔记二",
                "达人": "达人乙",
                "消费": 80,
                "展现量": 1000,
                "点击量": 50,
                "互动量": 20,
                "行业商品进店量30日": 8,
                "行业商品GMV（30日)": 80,
            },
        ]
    ).to_csv(path, index=False, encoding="utf-8-sig")


def _write_match(path: Path) -> None:
    rows = [
        [None, "昕麟投放账号", None, None, None, None, "非昕麟投放账号", None, None],
        ["笔记ID", "昵称", "发布链接", "发布日期", None, None, "昵称", "发布连接", "笔记ID"],
        ["note-1", "达人甲", "https://www.xiaohongshu.com/explore/note-1?xsec_token=secret", "2026-09-01", None, None, None, None, None],
        [None, None, None, None, None, None, "达人乙", "https://www.xiaohongshu.com/explore/note-2?xsec_token=other", "note-2"],
    ]
    pd.DataFrame(rows).to_excel(path, sheet_name="笔记匹配表", index=False, header=False)


def test_review_endpoint_returns_canonical_links_and_quality(tmp_path):
    _write_creative(tmp_path / "日报.csv")
    _write_match(tmp_path / "匹配.xlsx")
    client = TestClient(create_app(input_dir=tmp_path, default_source=None))

    response = client.get("/api/review", params={"start": "2026-09-01", "end": "2026-09-02"})

    assert response.status_code == 200
    payload = response.json()
    assert payload["project_id"] == "mgs-local"
    assert payload["summary"]["spend"] == 200
    assert payload["quality"]["note_link_coverage"] == 1
    urls = {item["note_url"] for item in payload["note_breakdown"]}
    assert urls == {
        "https://www.xiaohongshu.com/explore/note-1",
        "https://www.xiaohongshu.com/explore/note-2",
    }
    assert all("xsec_token" not in url for url in urls)


def test_review_export_contains_all_review_sheets(tmp_path):
    _write_creative(tmp_path / "日报.csv")
    _write_match(tmp_path / "匹配.xlsx")
    client = TestClient(create_app(input_dir=tmp_path, default_source=None))

    response = client.get("/api/export/review-xlsx", params={"start": "2026-09-01", "end": "2026-09-02"})

    assert response.status_code == 200
    output = tmp_path / "review.xlsx"
    output.write_bytes(response.content)
    import openpyxl

    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    assert workbook.sheetnames == [
        "复盘首页",
        "复盘结论",
        "Goodcase",
        "Badcase",
        "知识库命中",
        "复盘趋势",
        "点位复盘",
        "定向复盘",
        "关键词复盘",
        "达人复盘",
        "达人点评",
        "达人最佳定向",
        "笔记链接",
        "复盘数据质量",
    ]


def test_review_notes_are_saved_locally_without_touching_action_log(tmp_path):
    _write_creative(tmp_path / "日报.csv")
    client = TestClient(create_app(input_dir=tmp_path, default_source=None))

    response = client.post(
        "/api/review/notes",
        json={
            "note_id": "note-1",
            "creator_name": "达人甲",
            "note_url": "https://www.xiaohongshu.com/explore/note-1?xsec_token=secret",
            "status": "已点评",
            "comment": "钩子清晰，需补产品证据",
        },
    )

    assert response.status_code == 200
    assert response.json()["record"]["note_url"] == "https://www.xiaohongshu.com/explore/note-1"
    notes_path = tmp_path / "mgs_review_notes.jsonl"
    assert notes_path.exists()
    saved = json.loads(notes_path.read_text(encoding="utf-8").splitlines()[0])
    assert saved["project_id"] == "mgs-local"
    assert not (tmp_path / "mgs_action_log.jsonl").exists()


def test_review_notes_are_merged_into_review_and_export(tmp_path):
    _write_creative(tmp_path / "日报.csv")
    client = TestClient(create_app(input_dir=tmp_path, default_source=None))
    saved = client.post(
        "/api/review/notes",
        json={
            "note_id": "note-1",
            "note_url": "https://www.xiaohongshu.com/explore/note-1?xsec_token=secret",
            "status": "已点评",
            "comment": "开头钩子具体，建议复用",
            "reuse_suggestion": "改成搜索素材",
        },
    )
    assert saved.status_code == 200

    payload = client.get("/api/review", params={"start": "2026-09-01", "end": "2026-09-02"}).json()
    note = next(item for item in payload["note_breakdown"] if item["note_id"] == "note-1")
    assert note["review_status"] == "已点评"
    assert note["review_comment"] == "开头钩子具体，建议复用"
    assert note["reuse_suggestion"] == "改成搜索素材"

    import openpyxl
    output = tmp_path / "review.xlsx"
    response = client.get("/api/export/review-xlsx", params={"start": "2026-09-01", "end": "2026-09-02"})
    output.write_bytes(response.content)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    review_rows = list(workbook["达人点评"].iter_rows(values_only=True))
    assert any("已点评" in row for row in review_rows)
    assert any("开头钩子具体，建议复用" in row for row in review_rows)


def test_review_quality_sheet_includes_ingest_issues_and_unmatched_notes(tmp_path):
    _write_creative(tmp_path / "日报.csv")
    client = TestClient(create_app(input_dir=tmp_path, default_source=None))
    response = client.get("/api/export/review-xlsx", params={"start": "2026-09-01", "end": "2026-09-02"})
    assert response.status_code == 200

    import openpyxl
    output = tmp_path / "review-quality.xlsx"
    output.write_bytes(response.content)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    values = [str(value or "") for row in workbook["复盘数据质量"].iter_rows(values_only=True) for value in row]
    assert any("未匹配" in value for value in values)
