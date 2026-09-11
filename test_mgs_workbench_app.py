import io
import json
import os
import re
from pathlib import Path
from urllib.parse import unquote

from fastapi.testclient import TestClient

from mgs_workbench_app import create_app


def csv_bytes(rows: list[list[object]]) -> bytes:
    lines = [",".join(str(value) for value in row) for row in rows]
    return ("\ufeff" + "\n".join(lines)).encode("utf-8")


def test_health_and_default_state(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)

    response = client.get("/api/health")

    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_default_app_uses_example_registry_when_local_registry_is_missing(tmp_path, monkeypatch):
    import mgs_workbench_app as workbench

    example = tmp_path / "projects.example.json"
    example.write_text(
        json.dumps(
            {
                "projects": [
                    {
                        "project_id": "example-project",
                        "name": "示例项目",
                        "display_name": "示例项目",
                        "root_dir": str(tmp_path / "project"),
                        "input_dir": str(tmp_path / "project" / "输入"),
                        "default_source": "",
                        "kpi": {},
                        "action_rules": {},
                        "metric_contract": {},
                        "attribution_windows": {},
                        "modules": {"投放分析": True},
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(workbench, "DEFAULT_REGISTRY_PATH", tmp_path / "missing-projects.json")
    monkeypatch.setattr(workbench, "DEFAULT_EXAMPLE_REGISTRY_PATH", example)

    app = workbench.create_app()

    assert app.state.project.project_id == "example-project"


def test_file_signature_changes_only_after_input_file_changes(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    source = tmp_path / "日报.csv"
    source.write_bytes(
        csv_bytes(
            [
                ["时间", "计划名称", "消费"],
                ["2026-09-01", "计划A", "100"],
            ]
        )
    )

    first = client.get("/api/files").json()["signature"]
    same = client.get("/api/files").json()["signature"]
    source.write_bytes(source.read_bytes() + b"\n")
    changed = client.get("/api/files").json()["signature"]

    assert first == same
    assert changed != first


def test_latest_project_workbook_replaces_stale_input_copy(tmp_path):
    project_dir = tmp_path / "MGS投流"
    input_dir = project_dir / "数据输入"
    input_dir.mkdir(parents=True)
    stale_copy = input_dir / "MGS投流CID优化表-9.4.xlsx"
    independent_report = input_dir / "关键词-投放数据.csv"
    older_main = project_dir / "MGS投流CID优化表-9.6.xlsx"
    latest_main = project_dir / "MGS投流CID优化表-9.7.xlsx"
    for path in (stale_copy, independent_report, older_main, latest_main):
        path.write_bytes(b"test")
    os.utime(older_main, (100, 100))
    os.utime(latest_main, (300, 300))

    app = create_app(input_dir=input_dir, default_source=project_dir)
    files = TestClient(app).get("/api/files").json()["files"]

    paths = {item["path"] for item in files}
    assert str(latest_main) in paths
    assert str(independent_report) in paths
    assert str(stale_copy) not in paths
    assert str(older_main) not in paths


def test_explicit_primary_file_wins_over_same_named_input_copy(tmp_path):
    input_dir = tmp_path / "数据输入"
    project_dir = tmp_path / "MGS投流"
    input_dir.mkdir()
    project_dir.mkdir()
    filename = "MGS投流CID优化表-9.7.xlsx"
    stale_copy = input_dir / filename
    primary = project_dir / filename
    stale_copy.write_bytes(b"old")
    primary.write_bytes(b"new")

    app = create_app(input_dir=input_dir, default_source=primary)
    files = TestClient(app).get("/api/files").json()["files"]

    assert [item["path"] for item in files] == [str(primary)]


def test_newer_versioned_workbook_in_input_dir_wins_over_stale_project_copy(tmp_path):
    """用户把当天新主表放进数据输入时必须被读到，不能退回上级目录旧版本。

    版本判定看文件名（9.11 > 9.8），不看 mtime——旧拷贝可能只是后来被碰过。
    """
    project_dir = tmp_path / "MGS投流"
    input_dir = project_dir / "数据输入"
    project_dir.mkdir()
    input_dir.mkdir()
    older_main = project_dir / "MGS投流CID优化表-9.8.xlsx"
    newer_main = input_dir / "MGS投流CID优化表-9.11.xlsx"
    older_main.write_bytes(b"old")
    newer_main.write_bytes(b"new")
    # 故意把旧主表的 mtime 设得更晚，锁死「版本号优先于 mtime」
    os.utime(older_main, (300, 300))
    os.utime(newer_main, (100, 100))

    app = create_app(input_dir=input_dir, default_source=project_dir)
    files = TestClient(app).get("/api/files").json()["files"]

    assert [Path(item["path"]).name for item in files if item["path"].endswith(".xlsx")] == [
        "MGS投流CID优化表-9.11.xlsx"
    ]


def test_upload_then_analyze_returns_kpis_actions_and_issues(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "投放类型", "计划名称", "笔记ID", "消费", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
            ["2026-09-01", "信息流", "计划A", "note-1", "1000", "100", "20", "3000"],
            ["2026-09-02", "全站", "计划B", "note-2", "1000", "100", "5", "100"],
        ]
    )

    upload = client.post(
        "/api/upload",
        files={"files": ("创意-投放数据.csv", io.BytesIO(content), "text/csv")},
    )
    assert upload.status_code == 200
    assert upload.json()["accepted"] == 1

    analysis = client.post(
        "/api/analyze",
        json={"start": "2026-09-01", "end": "2026-09-02", "point": "全部"},
    )
    assert analysis.status_code == 200
    payload = analysis.json()
    assert payload["kpis"]["spend"] == 2000
    assert payload["kpis"]["roi"] == 1.55
    assert payload["actions"]
    assert payload["points"]
    assert payload["meta"]["file_count"] == 1


def test_report_and_xlsx_exports_use_latest_analysis(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "投放类型", "计划名称", "笔记ID", "消费", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
            ["2026-09-01", "信息流", "计划A", "note-1", "1000", "100", "20", "3000"],
        ]
    )
    client.post("/api/upload", files={"files": ("日报.csv", io.BytesIO(content), "text/csv")})
    client.post("/api/analyze", json={"start": "2026-09-01", "end": "2026-09-01", "point": "全部"})

    report = client.get("/api/export/report")
    workbook = client.get("/api/export/xlsx")

    assert report.status_code == 200
    assert "投放动作" in report.text
    assert workbook.status_code == 200
    assert workbook.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(workbook.content) > 1000


def test_upload_accepts_manual_report_kind_and_persists_manifest(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["日期", "人群包名称", "精准定向", "消费", "店铺访问量(15日)", "店铺成交ROI(30日)"],
            ["2026-09-01", "高意向人群", "高意向人群", "100", "4", "2"],
        ]
    )

    upload = client.post(
        "/api/upload",
        data={"kinds": json.dumps({"0": "targeting"})},
        files={"files": ("没有定向关键词的报表.csv", io.BytesIO(content), "text/csv")},
    )

    assert upload.status_code == 200
    assert upload.json()["files"][0]["kind"] == "targeting"
    manifest = json.loads((tmp_path / ".mgs_workbench_manifest.json").read_text(encoding="utf-8"))
    assert manifest["没有定向关键词的报表.csv"] == "targeting"


def test_file_kind_endpoint_updates_existing_upload(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "计划名称", "关键词", "消费"],
            ["2026-09-01", "搜索计划", "面膜", "100"],
        ]
    )
    client.post("/api/upload", files={"files": ("报表.csv", io.BytesIO(content), "text/csv")})

    response = client.post("/api/files/kinds", json={"kinds": {"报表.csv": "keyword"}})

    assert response.status_code == 200
    assert response.json()["files"][0]["kind"] == "keyword"


def test_file_kind_endpoint_clears_manual_kind_for_auto_detection(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "计划名称", "关键词", "消费"],
            ["2026-09-01", "搜索计划", "面膜", "100"],
        ]
    )
    client.post(
        "/api/upload",
        data={"kinds": json.dumps({"0": "targeting"})},
        files={"files": ("报表.csv", io.BytesIO(content), "text/csv")},
    )

    response = client.post("/api/files/kinds", json={"kinds": {"报表.csv": ""}})

    assert response.status_code == 200
    assert response.json()["files"][0]["kind"] == ""
    manifest = json.loads((tmp_path / ".mgs_workbench_manifest.json").read_text(encoding="utf-8"))
    assert "报表.csv" not in manifest


def test_compare_can_be_disabled(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "投放类型", "计划名称", "笔记ID", "消费", "点击量", "行业商品进店量30日", "行业商品GMV（30日)"],
            ["2026-08-31", "全站", "旧计划", "n-old", "100", "10", "5", "300"],
            ["2026-09-01", "全站", "新计划", "n-new", "100", "10", "5", "100"],
        ]
    )
    client.post("/api/upload", files={"files": ("日报.csv", io.BytesIO(content), "text/csv")})

    response = client.post(
        "/api/analyze",
        json={"start": "2026-09-01", "end": "2026-09-01", "compare": False},
    )

    assert response.status_code == 200
    assert response.json()["comparison"] == {"start": "", "end": ""}
    assert response.json()["baseline_kpis"] == {}


def test_analyze_keeps_all_point_options_after_filtering_one_point(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "投放类型", "计划名称", "笔记ID", "消费", "行业商品进店量30日", "行业商品GMV（30日)"],
            ["2026-09-01", "全站", "计划A", "n-a", "100", "5", "300"],
            ["2026-09-01", "信息流", "计划B", "n-b", "100", "5", "300"],
        ]
    )
    client.post("/api/upload", files={"files": ("日报.csv", io.BytesIO(content), "text/csv")})

    response = client.post(
        "/api/analyze",
        json={"start": "2026-09-01", "end": "2026-09-01", "point": "全站", "compare": False},
    )

    assert response.status_code == 200
    assert set(response.json()["point_options"]) == {"全站", "信息流"}


def test_xlsx_export_includes_all_note_drilldown_rows(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    rows = [
        ["时间", "投放类型", "计划名称", "笔记ID", "创意名称", "消费", "行业商品进店量30日", "行业商品GMV（30日)"],
    ]
    for index in range(85):
        rows.append(["2026-09-01", "全站", f"计划-{index}", f"note-{index}", f"笔记-{index}", "100", "2", "300"])
    content = csv_bytes(rows)
    client.post("/api/upload", files={"files": ("创意.csv", io.BytesIO(content), "text/csv")})

    workbook_response = client.get("/api/export/xlsx")

    assert workbook_response.status_code == 200
    output = tmp_path / "export.xlsx"
    output.write_bytes(workbook_response.content)
    import openpyxl

    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    assert workbook["笔记动作"].max_row - 4 == 85


def test_open_folder_endpoint_uses_local_folder_only(tmp_path, monkeypatch):
    import mgs_workbench_app

    calls = []

    def fake_popen(arguments, **kwargs):
        calls.append((arguments, kwargs))

    monkeypatch.setattr(mgs_workbench_app.subprocess, "Popen", fake_popen)
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)

    response = client.post("/api/open-folder")

    assert response.status_code == 200
    assert response.json()["path"] == str(tmp_path)
    assert calls == [(["open", str(tmp_path)], {"stdout": mgs_workbench_app.subprocess.DEVNULL, "stderr": mgs_workbench_app.subprocess.DEVNULL})]


def test_auto_kind_is_reinferred_after_clearing_manifest(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes(
        [
            ["时间", "计划名称", "关键词", "消费"],
            ["2026-09-01", "搜索计划", "面膜", "100"],
        ]
    )
    client.post(
        "/api/upload",
        data={"kinds": json.dumps({"0": "targeting"})},
        files={"files": ("报表.csv", io.BytesIO(content), "text/csv")},
    )
    client.post("/api/files/kinds", json={"kinds": {"报表.csv": ""}})

    response = client.post("/api/analyze", json={"start": "2026-09-01", "end": "2026-09-01", "compare": False})

    assert response.status_code == 200
    assert response.json()["available_kinds"] == ["keyword"]


def test_project_switch_changes_data_directory_kpi_and_action_log(tmp_path):
    import json

    first_root = tmp_path / "first"
    second_root = tmp_path / "second"
    first_input = first_root / "输入"
    second_input = second_root / "输入"
    first_input.mkdir(parents=True)
    second_input.mkdir(parents=True)
    rows = [
        ["时间", "投放类型", "计划名称", "消费", "行业商品进店量30日", "行业商品GMV（30日)"],
        ["2026-09-01", "全站", "计划", "100", "10", "300"],
    ]
    first_input.joinpath("first.csv").write_bytes(csv_bytes(rows))
    second_input.joinpath("second.csv").write_bytes(csv_bytes(rows))
    registry_path = tmp_path / "projects.json"
    registry_path.write_text(
        json.dumps(
            {
                "projects": [
                    {
                        "project_id": "first-project",
                        "name": "第一项目",
                        "display_name": "第一商品",
                        "root_dir": str(first_root),
                        "input_dir": str(first_input),
                        "default_source": "",
                        "kpi": {"roi_target": 2},
                        "metric_contract": {"roi": "gmv / spend"},
                        "modules": {"投放分析": True},
                    },
                    {
                        "project_id": "second-project",
                        "name": "第二项目",
                        "display_name": "第二商品",
                        "root_dir": str(second_root),
                        "input_dir": str(second_input),
                        "default_source": "",
                        "kpi": {"roi_target": 5},
                        "metric_contract": {"roi": "gmv / spend"},
                        "modules": {"投放分析": True},
                    },
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    app = create_app(registry_path=registry_path)
    client = TestClient(app)

    projects = client.get("/api/projects").json()
    assert projects["current_project_id"] == "first-project"
    assert [item["project_id"] for item in projects["projects"]] == ["first-project", "second-project"]

    first_state = client.get("/api/state").json()
    assert first_state["project_id"] == "first-project"
    assert first_state["targets"]["roi_target"] == 2
    assert first_state["meta"]["input_dir"] == str(first_input)

    switched = client.post("/api/projects/select", json={"project_id": "second-project"})
    assert switched.status_code == 200
    second_state = client.get("/api/state").json()
    assert second_state["project_id"] == "second-project"
    assert second_state["targets"]["roi_target"] == 5
    assert second_state["meta"]["input_dir"] == str(second_input)

    client.post("/api/actions/log", json={"action": "观察", "entity": "第二计划"})
    assert not (first_input / "mgs_action_log.jsonl").exists()
    assert (second_input / "mgs_action_log.jsonl").exists()


def test_project_exports_use_current_project_name(tmp_path):
    import json

    project_root = tmp_path / "client"
    input_dir = project_root / "输入"
    input_dir.mkdir(parents=True)
    input_dir.joinpath("日报.csv").write_bytes(
        csv_bytes([
            ["时间", "投放类型", "计划名称", "消费", "行业商品进店量30日", "行业商品GMV（30日)"],
            ["2026-09-01", "全站", "计划", "100", "10", "300"],
        ])
    )
    registry_path = tmp_path / "projects.json"
    registry_path.write_text(
        json.dumps({
            "projects": [{
                "project_id": "client-project",
                "name": "客户A",
                "display_name": "客户A精华液",
                "root_dir": str(project_root),
                "input_dir": str(input_dir),
                "default_source": "",
                "kpi": {"roi_target": 1.2},
                "metric_contract": {"roi": "gmv / spend"},
                "attribution_windows": {"gmv_days": 7, "visits_days": 3},
                "modules": {"投放分析": True},
            }]
        }, ensure_ascii=False),
        encoding="utf-8",
    )

    client = TestClient(create_app(registry_path=registry_path))
    report = client.get("/api/export/report")
    workbook = client.get("/api/export/xlsx")

    report_disposition = unquote(report.headers["content-disposition"])
    workbook_disposition = unquote(workbook.headers["content-disposition"])
    assert "客户A投放日报" in report_disposition
    assert "客户A投放分析工作台" in workbook_disposition
    assert "店铺成交GMV（7日）" in report.text
    assert "店铺成交ROI（7日）" in report.text


def test_account_filter_separates_red_cat_and_red京(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes([
        ["时间", "投放类型", "计划名称", "消费", "行业商品进店量30日", "行业商品GMV（30日)", "下载账户"],
        ["2026-09-01", "全站", "红猫计划", "100", "10", "300", "红猫"],
        ["2026-09-01", "全站", "红京计划", "200", "20", "800", "红京"],
    ])
    uploaded = client.post("/api/upload", files={"files": ("日报.csv", io.BytesIO(content), "text/csv")})
    assert uploaded.status_code == 200

    all_accounts = client.get("/api/state").json()
    assert all_accounts["account_options"] == ["红京", "红猫"]
    red_cat = client.post("/api/analyze", json={"account_id": "红猫", "start": "2026-09-01", "end": "2026-09-01", "compare": False}).json()
    red_jing = client.post("/api/analyze", json={"account_id": "红京", "start": "2026-09-01", "end": "2026-09-01", "compare": False}).json()
    assert red_cat["kpis"]["spend"] == 100
    assert red_cat["kpis"]["gmv"] == 300
    assert red_jing["kpis"]["spend"] == 200
    assert red_jing["kpis"]["gmv"] == 800


def test_report_export_scopes_download_to_requested_account(tmp_path):
    app = create_app(input_dir=tmp_path, default_source=None)
    client = TestClient(app)
    content = csv_bytes([
        ["时间", "投放类型", "计划名称", "消费", "行业商品进店量30日", "行业商品GMV（30日)", "下载账户"],
        ["2026-09-01", "全站", "红猫计划", "100", "10", "300", "红猫"],
        ["2026-09-01", "全站", "红京计划", "200", "20", "800", "红京"],
    ])
    uploaded = client.post("/api/upload", files={"files": ("日报.csv", io.BytesIO(content), "text/csv")})
    assert uploaded.status_code == 200

    red_jing = client.get("/api/export/report", params={"account_id": "红京"})

    assert red_jing.status_code == 200
    assert "红京" in unquote(red_jing.headers["content-disposition"])
    assert "消耗：200.00" in red_jing.text
    assert "店铺成交GMV（30日）：800.00" in red_jing.text
    assert "消耗：100.00" not in red_jing.text


def test_report_page_download_preserves_selected_account():
    static_app = Path(__file__).parents[1] / "mgs_workbench" / "mgs_workbench_static" / "app.js"
    source = static_app.read_text(encoding="utf-8")
    index = static_app.with_name("index.html").read_text(encoding="utf-8")

    assert '$("#report-download").addEventListener("click", () => download(`/api/export/report${exportAccountQuery()}`));' in source
    assert 'reviewExportPath' in source and 'params.set("account_id", state.accountFilter)' in source
    assert 'reviewCreatorTargetingFilter' in source
    assert '$("#review-creator-targeting-select").addEventListener' in source
    assert 'allRows.filter((row) => String(row.creator_name || "") === state.reviewCreatorTargetingFilter)' in source
    assert 'id="review-creator-targeting-select"' in index
    # 版本号每次改前端都会 bump；断言形状而不是具体数字，避免测试跟着改。
    assert re.search(r"app\.js\?v=\d{8}-\d+", index), "index.html 必须给 app.js 带缓存版本号"
    assert re.search(r"styles\.css\?v=\d{8}-\d+", index), "index.html 必须给 styles.css 带缓存版本号"


def test_default_analysis_window_starts_at_month_beginning_and_ends_yesterday(tmp_path):
    import mgs_workbench_app as workbench
    import pandas as pd

    yesterday = workbench.date.today() - workbench.timedelta(days=1)
    month_start = yesterday.replace(day=1)
    frame = pd.DataFrame({"date": pd.to_datetime([month_start, yesterday]), "spend": [1, 2]})

    assert workbench._default_range({"creative": frame}) == (month_start.isoformat(), yesterday.isoformat())


def test_write_table_supports_columns_beyond_z():
    from openpyxl import Workbook

    import mgs_workbench_app as workbench

    sheet = Workbook().active
    headers = [f"列{index}" for index in range(1, 31)]
    workbench._write_table(sheet, headers, [list(range(1, 31))])

    refs = [table.ref for table in sheet.tables.values()]
    assert refs == ["A4:AD5"]
